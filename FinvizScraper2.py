
import requests
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import load_workbook
from io import StringIO

workbook = load_workbook('FinvizData2.xlsx')
sheet = workbook['Sheet1']

def findStartPoint():
    index = 1
    #if this loop is true, the current cell has a value inside of it.
    while(sheet.cell(row = index, column = 1).value != None):
        index += 1
    return index


def putDataIntoSheet(df):
    rowIndex = findStartPoint() - 1
    
    df = df.iloc[:, 1:]
    for ticker in range(len(df)):
        columnIndex = 1
        rowIndex += 1
        #find the row corresponding to the ticker index
        tickerDetails = df.iloc[ticker]
        for tickerDetail in tickerDetails:
            sheet.cell(row = rowIndex, column = columnIndex).value = tickerDetail
            columnIndex += 1

def getTotalPages(soup):
    try:
        paginationTags = soup.find(class_="body-table screener_pagination").find_all('a')
    except Exception as e:
        print("There are no results that match the selected criteria.  Please try again.")
        return
    
    totalPages = (len(paginationTags) - 1) #excluding the arrow
    if(totalPages == 0): #there is only one page
        return 1
    else:
        return totalPages




def getWebpage():
    url = "https://finviz.com/screener.ashx?v=121&f=fa_curratio_o1,fa_debteq_low,fa_pe_low,fa_peg_low,fa_quickratio_o1.5,fa_roe_pos&ft=4"
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36(KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"}
    
    #connect to our page
    response = requests.get(url, headers=headers)
    #read our response
    soup = BeautifulSoup(response.content, 'html.parser')

    numPages = getTotalPages(soup)

    
    
    tickerNumber = 1
    for numPage in range(numPages):
        #visit each page and convert into pandas data
        response = requests.get(url+f"&r={tickerNumber}", headers=headers)
        soup = BeautifulSoup(response.content, 'html.parser')
        table = soup.find('table', class_="styled-table-new is-rounded is-tabular-nums w-full screener_table")
        table_html = StringIO(str(table))
        pdData = pd.read_html(table_html)
        tickerNumber += 20
        putDataIntoSheet(pdData[0])

getWebpage()
workbook.save('FinvizData2.xlsx')




    

    



